"""Payroll and attendance workbooks (sensitive doc types).

They imitate exports from an HR/payroll system: a title block, a summary sheet with one
row per month, and one register sheet per month listing every worker with name, NIC and
bank account. Values are written as numbers, not formulas, because (1) that is what
payroll software exports and (2) openpyxl formulas have no cached values, so pandas
would read them back as empty.

All people, NIC numbers and bank accounts are randomly generated.
"""

from __future__ import annotations

import random
from dataclasses import dataclass
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .common import Factory, Month

FIRST_F = [
    "Nadeesha",
    "Dilani",
    "Chamari",
    "Sanduni",
    "Kumari",
    "Iresha",
    "Thilini",
    "Malsha",
    "Nirosha",
    "Hasini",
    "Sewwandi",
    "Anusha",
    "Priyanka",
    "Shanika",
    "Dilrukshi",
    "Kavindi",
    "Tharushi",
    "Lakshika",
    "Nilmini",
    "Samanthi",
    "Kalaivani",
    "Tharshini",
    "Sumathy",
    "Fathima",
    "Rizna",
    "Shafna",
    "Vasuki",
    "Mathivathani",
    "Ishara",
    "Gayani",
]
FIRST_M = [
    "Kasun",
    "Nuwan",
    "Chaminda",
    "Ruwan",
    "Saman",
    "Tharindu",
    "Dinesh",
    "Lahiru",
    "Pradeep",
    "Asanka",
    "Sujith",
    "Mohamed",
    "Rifkhan",
    "Kumaran",
    "Suresh",
    "Janaka",
    "Isuru",
    "Dulaj",
]
SURNAMES = [
    "Perera",
    "Fernando",
    "Silva",
    "Jayasinghe",
    "Wickramasinghe",
    "Bandara",
    "Herath",
    "Rathnayake",
    "Dissanayake",
    "Gunawardena",
    "Karunaratne",
    "Kumari",
    "Madushani",
    "Senanayake",
    "Weerasinghe",
    "Ekanayake",
    "Rajapaksha",
    "Samarakoon",
    "Liyanage",
    "Pathirana",
    "Sivakumar",
    "Thevarajah",
    "Nazeer",
    "Ranasinghe",
    "Abeysekara",
    "Hettiarachchi",
]
DEPTS = [
    ("Sewing", "Sewing Machine Operator", 0.62),
    ("Cutting", "Cutter", 0.08),
    ("Quality", "Quality Checker", 0.08),
    ("Finishing", "Ironer / Packer", 0.12),
    ("Stores", "Store Assistant", 0.04),
    ("Maintenance", "Mechanic", 0.03),
    ("Administration", "Clerk", 0.03),
]
BANKS = ["People's Bank", "Bank of Ceylon", "Sampath Bank", "Commercial Bank", "HNB"]

HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=9)
BODY_FONT = Font(name="Arial", size=9)
TITLE_FONT = Font(name="Arial", bold=True, size=12)
THIN = Side(style="thin", color="B7C0CC")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)


@dataclass
class Worker:
    epf_no: str
    name: str
    nic: str
    dept: str
    designation: str
    basic_2025: int
    bank: str
    account: str
    joined: int  # index of first month employed
    left: int  # index of last month employed


def _nic(rng: random.Random, female: bool) -> str:
    """Format-valid NIC: old 9 digits + V, or new 12 digits (+500 day offset for women)."""
    year = rng.randint(1975, 2005)
    day = rng.randint(1, 365) + (500 if female else 0)
    if year < 2000 and rng.random() < 0.55:
        return f"{year % 100:02d}{day:03d}{rng.randint(0, 999):03d}{rng.randint(0, 9)}V"
    return f"{year}{day:03d}0{rng.randint(0, 999):03d}{rng.randint(0, 9)}"


def build_workforce(rng: random.Random, size: int, months: int, epf_prefix: str) -> list[Worker]:
    people: list[Worker] = []
    weights = [d[2] for d in DEPTS]
    for i in range(size):
        female = rng.random() < 0.82
        first = rng.choice(FIRST_F if female else FIRST_M)
        dept, desig, _ = rng.choices(DEPTS, weights)[0]
        base = {
            "Sewing": 29_500,
            "Cutting": 31_000,
            "Quality": 31_500,
            "Finishing": 29_500,
            "Stores": 30_500,
            "Maintenance": 36_000,
            "Administration": 34_000,
        }[dept]
        joined = 0 if rng.random() < 0.9 else rng.randint(1, months - 1)
        left = months - 1 if rng.random() < 0.94 else rng.randint(joined, months - 1)
        people.append(
            Worker(
                epf_no=f"{epf_prefix}{1000 + i:05d}",
                name=f"{first} {rng.choice(SURNAMES)}",
                nic=_nic(rng, female),
                dept=dept,
                designation=desig,
                basic_2025=base + rng.choice([0, 0, 0, 500, 1000, 1500, 2500]),
                bank=rng.choice(BANKS),
                account=f"{rng.randint(10**11, 10**12 - 1)}",
                joined=joined,
                left=left,
            )
        )
    return people


def _basic_for(w: Worker, m: Month) -> int:
    """Pay rise from January 2026 when the national minimum moved 27,000 -> 30,000 LKR."""
    return w.basic_2025 + (2_000 if (m.year, m.month) >= (2026, 1) else 0)


def _title(ws: Worksheet, f: Factory, report: str, generated: date, ncols: int) -> None:
    ws["A1"] = f.name
    ws["A1"].font = TITLE_FONT
    ws["A2"] = report
    ws["A2"].font = Font(name="Arial", bold=True, size=10)
    ws["A3"] = f"Exported from PayMaster HR 7.2 on {generated.strftime('%d/%m/%Y')} | Confidential"
    ws["A3"].font = Font(name="Arial", italic=True, size=8, color="666666")
    for r in (1, 2, 3):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)


def _table(
    ws: Worksheet,
    header_row: int,
    headers: list[str],
    rows: list[list[object]],
    formats: dict[int, str],
    widths: list[int],
) -> None:
    for c, h in enumerate(headers, 1):
        cell = ws.cell(header_row, c, h)
        cell.font, cell.fill, cell.border = HEAD_FONT, HEAD_FILL, BOX
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 30
    for r, row in enumerate(rows, header_row + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v)
            cell.font, cell.border = BODY_FONT, BOX
            if c in formats:
                cell.number_format = formats[c]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(header_row + 1, 1)


def payroll_workbook(
    f: Factory,
    workers: list[Worker],
    months: list[Month],
    skip: set[str],
    rng: random.Random,
    generated: date,
) -> tuple[Workbook, list[dict[str, object]]]:
    wb = Workbook()
    summary = _first_sheet(wb)
    summary.title = "Summary"
    series: list[dict[str, object]] = []
    for idx, m in enumerate(months):
        if m.period in skip:
            continue  # planted gap: this month was never exported
        ws = wb.create_sheet(m.first_day.strftime("%b-%Y"))
        rows: list[list[object]] = []
        ot_total, gross_total, lowest = 0.0, 0.0, 10**9
        for w in workers:
            if not (w.joined <= idx <= w.left):
                continue
            basic = _basic_for(w, m)
            att_allow = 2_500 if rng.random() < 0.86 else 0
            ot_hours = round(rng.choice([0, 0, 8, 12, 16, 20, 24, 30, 36]) + rng.random() * 3, 1)
            ot_rate = round(basic / 200 * 1.5, 2)  # monthly basic / 200 x 1.5
            ot_pay = round(ot_hours * ot_rate, 2)
            gross = round(basic + att_allow + ot_pay, 2)
            epf = round((basic + att_allow) * 0.08, 2)
            net = round(gross - epf, 2)
            rows.append(
                [
                    w.epf_no,
                    w.name,
                    w.nic,
                    w.dept,
                    w.designation,
                    basic,
                    att_allow,
                    ot_hours,
                    ot_rate,
                    ot_pay,
                    gross,
                    epf,
                    net,
                    w.bank,
                    w.account,
                ]
            )
            ot_total += ot_hours
            gross_total += gross
            lowest = min(lowest, basic)
        _title(ws, f, f"Monthly Payroll Register - {m.label}", generated, 15)
        _table(
            ws,
            5,
            [
                "EPF No",
                "Employee Name",
                "NIC No",
                "Department",
                "Designation",
                "Basic Salary (LKR)",
                "Attendance Allowance (LKR)",
                "OT Hours",
                "OT Rate (LKR/hr)",
                "OT Amount (LKR)",
                "Gross Earnings (LKR)",
                "EPF Employee 8% (LKR)",
                "Net Pay (LKR)",
                "Bank",
                "Account No",
            ],
            rows,
            {
                6: "#,##0.00",
                7: "#,##0.00",
                8: "0.0",
                9: "#,##0.00",
                10: "#,##0.00",
                11: "#,##0.00",
                12: "#,##0.00",
                13: "#,##0.00",
            },
            [11, 24, 15, 14, 22, 13, 13, 9, 11, 13, 14, 13, 13, 16, 15],
        )
        series.append(
            {
                "period": m.period,
                "worker_count": len(rows),
                "total_gross_wages": round(gross_total, 2),
                "lowest_basic_wage": lowest,
                "overtime_hours_total": round(ot_total, 1),
                "overtime_paid_flag": True,
            }
        )
    _title(summary, f, "Payroll Summary by Month (LKR)", generated, 7)
    _table(
        summary,
        5,
        [
            "Period",
            "Month",
            "Number of Workers",
            "Total Gross Wages (LKR)",
            "Lowest Basic Wage (LKR)",
            "Total Overtime Hours",
            "Overtime Paid",
        ],
        [
            [
                s["period"],
                _label(str(s["period"])),
                s["worker_count"],
                s["total_gross_wages"],
                s["lowest_basic_wage"],
                s["overtime_hours_total"],
                "Yes" if s["overtime_paid_flag"] else "No",
            ]
            for s in series
        ],
        {4: "#,##0.00", 5: "#,##0.00", 6: "#,##0.0"},
        [10, 16, 12, 18, 16, 14, 11],
    )
    return wb, series


def attendance_workbook(
    f: Factory, workers: list[Worker], months: list[Month], rng: random.Random, generated: date
) -> tuple[Workbook, list[dict[str, object]]]:
    wb = Workbook()
    summary = _first_sheet(wb)
    summary.title = "Summary"
    series: list[dict[str, object]] = []
    for idx, m in enumerate(months):
        ws = wb.create_sheet(m.first_day.strftime("%b-%Y"))
        weeks = 5 if m.last_day.day - (6 - m.first_day.weekday()) > 28 else 4
        rows: list[list[object]] = []
        all_hours: list[float] = []
        for w in workers:
            if not (w.joined <= idx <= w.left):
                continue
            hours = [
                round(
                    min(57.5, 45 + rng.choice([0, 0, 2, 4, 6, 8, 10, 12]) + rng.random() * 0.5), 1
                )
                for _ in range(weeks)
            ]
            if weeks == 5:
                hours[-1] = round(hours[-1] * rng.uniform(0.3, 0.6), 1)  # partial last week
            full_weeks = hours[:4]
            all_hours.extend(full_weeks)
            rows.append(
                [w.epf_no, w.name, w.dept, *hours, max(full_weeks), round(sum(full_weeks) / 4, 1)]
            )
        headers = (
            ["EPF No", "Employee Name", "Department"]
            + [f"Week {i + 1} Hours" for i in range(weeks)]
            + ["Max Weekly Hours", "Avg Weekly Hours"]
        )
        _title(ws, f, f"Attendance and Working Hours Register - {m.label}", generated, len(headers))
        _table(ws, 5, headers, rows, {}, [11, 24, 14] + [10] * weeks + [11, 11])
        series.append(
            {
                "period": m.period,
                "max_weekly_hours": max(all_hours),
                "average_weekly_hours": round(sum(all_hours) / len(all_hours), 1),
            }
        )
    _title(summary, f, "Working Hours Summary by Month (full weeks only)", generated, 4)
    _table(
        summary,
        5,
        ["Period", "Month", "Max Weekly Hours", "Average Weekly Hours"],
        [
            [
                s["period"],
                _label(str(s["period"])),
                s["max_weekly_hours"],
                s["average_weekly_hours"],
            ]
            for s in series
        ],
        {3: "0.0", 4: "0.0"},
        [10, 16, 14, 16],
    )
    return wb, series


def _first_sheet(wb: Workbook) -> Worksheet:
    ws = wb.active
    if not isinstance(ws, Worksheet):
        raise RuntimeError("new workbook has no active worksheet")
    return ws


def _label(period: str) -> str:
    y, m = period.split("-")
    return Month(int(y), int(m)).label
